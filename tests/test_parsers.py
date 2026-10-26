import pytest
from services.cv_parser_service import CVParserService
from services.jd_parser_service import _parse_skill_list_to_dict

@pytest.fixture
def parser():
    return CVParserService()

def test_find_full_name_filters_headers(parser):
    cv_text = """
    CURRICULUM VITAE
    SENIOR PYTHON DEVELOPER
    Nguyen Van An
    Email: an.nguyen@gmail.com
    Phone: 0987654321
    Summary: 5 years experience in building web apps.
    """
    name = parser._find_full_name(cv_text)
    assert name == "Nguyen Van An"

def test_find_years_experience_from_keyword(parser):
    text1 = "I have 3 years experience as a software engineer."
    assert parser._find_years_experience(text1) == 3.0

def test_find_years_experience_from_date_range(parser):
    text2 = "Backend Developer at FPT Software (2021 - 2024)"
    assert parser._find_years_experience(text2) >= 3.0

def test_extract_emails(parser):
    text = "Contact me at candidate@example.com or backup_email123@work.vn"
    emails = parser._extract_emails(text)
    assert "candidate@example.com" in emails
    assert "backup_email123@work.vn" in emails

def test_extract_foreign_languages(parser):
    text = "Languages: English (IELTS 7.5), Japanese (N3), Vietnamese Native."
    langs = parser._extract_foreign_languages(text)
    assert any("English" in l for l in langs)
    assert any("Japanese" in l for l in langs)

def test_alias_normalization():
    raw_list = ["Python (Senior)", "K8s (Basic)", "Postgres (Good)", "JS (Mid)"]
    parsed = _parse_skill_list_to_dict(raw_list)
    assert "Python" in parsed and parsed["Python"] == 3
    assert "Kubernetes" in parsed and parsed["Kubernetes"] == 1
    assert "Postgresql" in parsed and parsed["Postgresql"] == 2
    assert "Javascript" in parsed and parsed["Javascript"] == 2


def test_experience_ignores_age(parser):
    text = "Born 1998. 25 years old.\nBackend Dev at FPT (2020 - 2023)"
    assert parser._find_years_experience(text) == 3.0


def test_experience_merges_overlapping_ranges(parser):
    text = "Company A (2020 - 2023)\nCompany B (2023 to 2025)"
    assert parser._find_years_experience(text) == 5.0


def test_experience_sums_disjoint_ranges(parser):
    text = "Startup (2019 - 2020)\nAgency (2022 - 2024)"
    assert parser._find_years_experience(text) == 3.0


def test_experience_skips_education_ranges(parser):
    text = "Education\n2016 - 2020 University of Science\nExperience\nDev at Corp (2020 - 2022)"
    assert parser._find_years_experience(text) == 2.0


def test_digital_pdf_is_not_ocr_scrubbed(tmp_path, parser):
    """Text số phải giữ nguyên: clean_ocr_errors_smart sẽ biến "Django4" thành "Djang04"."""
    import fitz
    pdf_path = tmp_path / "cv_digital.pdf"
    doc = fitz.open()
    page = doc.new_page()
    lines = ["Nguyen Van An", "Python Django4 REST API", "Contact: 1l111l11l1", *[f"line {i} of filler text here" for i in range(40)]]
    page.insert_textbox(fitz.Rect(20, 20, 550, 750), "\n".join(lines), fontsize=9)
    doc.save(pdf_path)
    doc.close()

    text, error, mode = parser.parse_cv_document(str(pdf_path))
    assert mode == "digital_pdf" and error is None
    assert "Django4" in text
    assert "1l111l11l1" in text


def test_docx_tables_are_extracted(tmp_path, parser):
    import docx
    docx_path = tmp_path / "cv_table.docx"
    document = docx.Document()
    document.add_paragraph("Backend Developer")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Email"
    table.rows[0].cells[1].text = "an.nguyen@example.com"
    table.rows[1].cells[0].text = "Skill"
    table.rows[1].cells[1].text = "Python FastAPI"
    document.save(str(docx_path))

    text, error, mode = parser.parse_cv_document(str(docx_path))
    assert mode == "digital_docx" and error is None
    assert "an.nguyen@example.com" in text
    assert "Python FastAPI" in text


