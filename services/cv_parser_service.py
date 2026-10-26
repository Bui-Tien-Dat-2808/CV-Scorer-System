import os
import re
import cv2
import numpy as np
import fitz  # PyMuPDF
from datetime import datetime
from typing import Tuple, Optional, Dict, Any, List

# --- Import Models & Utils ---
from models.ocr_model import OCRModel
from models.embedding_model import EmbeddingModel
from utils.image_processing import preprocess_column, split_cv_columns
from utils.text_cleaner import clean_ocr_errors_smart, clean_text
from utils.file_loader import load_file_as_images
from core.constants import (
    EDUCATION_LEVELS, LANGUAGES, FRAMEWORKS, DATABASES,
    AI_ML_SKILLS, DEVOPS_TOOLS, SKILL_ALIASES, NAME_IGNORE_KEYWORDS
)
from core.logger import logger
from schemas.models import CandidateProfile

# Tên ngoại ngữ kèm biến thể tiếng Việt/tiếng Anh thường gặp trong CV
LANGUAGE_SYNONYMS = {
    'English': [r'\benglish\b', r'\btiếng\s*anh\b', r'\banh\s*ngữ\b', r'\banh\s*văn\b'],
    'Japanese': [r'\bjapanese\b', r'\btiếng\s*nhật\b', r'\bnhật\s*ngữ\b'],
    'Korean': [r'\bkorean\b', r'\btiếng\s*hàn\b', r'\bhàn\s*ngữ\b'],
    'Chinese': [r'\bchinese\b', r'\bmandarin\b', r'\btiếng\s*trung\b', r'\btiếng\s*hoa\b', r'\btrung\s*ngữ\b'],
    'French': [r'\bfrench\b', r'\btiếng\s*pháp\b', r'\bpháp\s*ngữ\b'],
    'German': [r'\bgerman\b', r'\btiếng\s*đức\b', r'\bđức\s*ngữ\b'],
    'Spanish': [r'\bspanish\b', r'\btiếng\s*tây\s*ban\s*nha\b'],
    'Vietnamese': [r'\bvietnamese\b', r'\btiếng\s*việt\b'],
}

LANGUAGE_LEVEL_PATTERN = (
    r'(IELTS\s*[\d.]{1,4}|TOEIC\s*\d{3,4}|TOEFL(?:\s*iBT)?\s*[\d.]{1,3}'
    r'|JLPT\s*N?[1-5]|N[1-5]|HSK\s*[1-6]|TOPIK\s*[1-6]|DELF\s*[A-C][12]|Goethe[-\s]?[A-C][12]'
    r'|Native|Fluent|Proficient|Advanced|Intermediate|Basic|Bản ngữ|Thành thạo|Khá|Trung bình)'
)

# Dòng chứa các từ khóa này thuộc khối học vấn: khoảng năm đi kèm là năm học, không phải năm đi làm.
EDUCATION_LINE_KEYWORDS = (
    'university', 'college', 'institute', 'school', 'gpa', 'degree', 'major',
    'đại học', 'cao đẳng', 'trường', 'tốt nghiệp', 'cử nhân', 'kỹ sư', 'thạc sĩ',
)

# Nhận diện khoảng thời gian: "2021 - 2024", "2022/2025", "2020 đến nay", "2019 to present"
YEAR_RANGE_PATTERN = re.compile(
    r'\b(19\d{2}|20\d{2})\s*'
    r'(?:[-–—~]|/|\bto\b|\btill\b|\buntil\b|\bđến\b|\btới\b)'
    r'\s*(19\d{2}|20\d{2}|present|nay|now|hiện\s*tại|this\s*time)\b',
    re.IGNORECASE,
)

class CVParserService:
    def __init__(self):
        self._ocr = None
        self._embedder = None

    @property
    def ocr(self):
        """Lazy loading OCR Model chỉ khi gặp file scan/ảnh cần OCR."""
        if self._ocr is None:
            self._ocr = OCRModel.get_instance()
        return self._ocr

    @property
    def embedder(self):
        """Lazy loading embedding model để không làm chậm lúc khởi động."""
        if self._embedder is None:
            self._embedder = EmbeddingModel.get_instance()
        return self._embedder

    # =========================================================================
    # 1. HYBRID EXTRACTION ENGINE (Fast-path Digital PDF + Fallback OCR)
    # =========================================================================

    def parse_cv_document(self, file_path: str) -> Tuple[str, Optional[str], str]:
        """
        Trích xuất văn bản từ CV theo cơ chế Hybrid:
        - Nếu là Digital PDF (có sẵn text layer): Trích xuất trong 0.05s với PyMuPDF.
        - Nếu là Scanned PDF hoặc Ảnh: Tự động kích hoạt PaddleOCR.
        - Nếu là DOCX: Đọc trực tiếp qua python-docx.

        Returns:
            (clean_text, error_message, extraction_mode)
        """
        if not os.path.exists(file_path):
            return "", "File không tồn tại.", "unknown"

        file_ext = os.path.splitext(file_path)[1].lower()

        # ---------------------------------------------------------------------
        # Case 1: FILE DOCX
        # ---------------------------------------------------------------------
        if file_ext == ".docx":
            try:
                import docx
                doc = docx.Document(file_path)
                parts = [p.text for p in doc.paragraphs if p.text.strip()]
                # Bàn trong CV (thông tin cá nhân, kỹ năng, kinh nghiệm) nằm ở table,
                # python-docx không trả nó qua doc.paragraphs.
                for table in doc.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
                full_text = "\n".join(parts)
                if full_text.strip():
                    logger.info("⚡ Fast-path: Đã trích xuất văn bản từ DOCX (cả paragraph và table).")
                    return clean_text(full_text), None, "digital_docx"
            except Exception as e:
                logger.warning(f"Lỗi đọc file docx: {e}. Thử fallback.")

        # ---------------------------------------------------------------------
        # Case 2: FILE PDF (Kiểm tra xem có text số không)
        # ---------------------------------------------------------------------
        if file_ext == ".pdf":
            try:
                doc = fitz.open(file_path)
                digital_text = ""
                for page in doc:
                    digital_text += page.get_text("text") + "\n"
                doc.close()

                # Đếm số từ trong text trích xuất được
                word_count = len(digital_text.strip().split())

                # Nếu file có trên 40 từ -> Đây là Digital PDF!
                if word_count >= 40:
                    logger.info(f"⚡ Fast-path: Phát hiện Digital PDF ({word_count} từ). Bỏ qua OCR hoàn toàn!")
                    # Văn bản số đã chính xác; clean_ocr_errors_smart chỉ dành cho OCR
                    # và sẽ làm hỏng chuỗi hợp lệ (vd "Django4" -> "Djang04").
                    return clean_text(digital_text), None, "digital_pdf"

                logger.info(f"🔍 PDF dạng scan/ảnh (chỉ có {word_count} từ text số). Chuyển sang PaddleOCR...")
            except Exception as e:
                logger.warning(f"Lỗi khi thử đọc digital PDF: {e}. Chuyển sang OCR...")

        # ---------------------------------------------------------------------
        # Case 3: FALLBACK SANG OCR (Cho PDF scan và file Ảnh thông thường)
        # ---------------------------------------------------------------------
        images = load_file_as_images(file_path)
        if not images:
            return "", "Không thể load file hoặc file rỗng.", "error"

        full_text = ""
        for i, img in enumerate(images):
            logger.info(f"OCR đang xử lý trang {i + 1}/{len(images)}...")
            try:
                page_text = self._process_single_page_numpy(img)
                if page_text:
                    full_text += page_text + "\n"
            except Exception as e:
                logger.error(f"Lỗi OCR trang {i + 1}: {e}")
                continue

        if not full_text.strip():
            return "", "Không trích xuất được nội dung CV (OCR không nhận diện được chữ).", "failed"

        clean_raw = clean_ocr_errors_smart(clean_text(full_text))
        return clean_raw, None, "ocr_fallback"

    def _ocr_image_data(self, image_data):
        """Hàm gọi PaddleOCR trên mảng ảnh numpy."""
        try:
            result = self.ocr.ocr(image_data)
        except Exception as e:
            logger.error(f"OCR engine failed: {e}")
            return ""

        if not result or not result[0]:
            return ""

        lines = []
        for line in result[0]:
            try:
                text = line[1][0]
                score = line[1][1]
                if score > 0.45:
                    lines.append(text)
            except Exception:
                continue

        return "\n".join(lines)

    def _process_single_page_numpy(self, image_numpy):
        """Xử lý chia cột (Layout Analysis) thông minh."""
        if len(image_numpy.shape) == 3:
            gray = cv2.cvtColor(image_numpy, cv2.COLOR_BGR2GRAY)
        else:
            gray = image_numpy

        height, width = gray.shape
        left_check_area = gray[:, :int(width * 0.40)]
        right_check_area = gray[:, int(width * 0.60):]

        mean_left = np.mean(left_check_area)
        mean_right = np.mean(right_check_area)

        DIFFERENCE_THRESHOLD = 75
        is_dark_left = (abs(mean_left - mean_right) > DIFFERENCE_THRESHOLD) and (mean_left < mean_right)
        is_dark_right = (abs(mean_left - mean_right) > DIFFERENCE_THRESHOLD) and (mean_right < mean_left)

        if is_dark_left:
            split = int(width * 0.42)
            left_txt = self._ocr_image_data(preprocess_column(gray[:, :split], True))
            right_txt = self._ocr_image_data(preprocess_column(gray[:, split:], False))
            return left_txt + "\n\n" + right_txt
        elif is_dark_right:
            split = int(width * 0.58)
            left_txt = self._ocr_image_data(preprocess_column(gray[:, :split], False))
            right_txt = self._ocr_image_data(preprocess_column(gray[:, split:], True))
            return left_txt + "\n\n" + right_txt
        else:
            processed = preprocess_column(gray, np.mean(gray) < 100)
            return self._ocr_image_data(processed)

    # =========================================================================
    # 2. INFORMATION EXTRACTION (Entity & Feature Engineering)
    # =========================================================================

    def extract_info(self, text: str, mode: str = "hybrid") -> Dict[str, Any]:
        """Trích xuất toàn diện các thực thể từ văn bản CV."""
        return {
            "full_name": self._find_full_name(text),
            "email": self._extract_emails(text),
            "phone_number": self._find_phone_number(text),
            "years_experience": self._find_years_experience(text),
            "education_level": self._extract_education_level(text),
            "skills_detected": self._extract_all_skills_passive(text),
            "languages_detected": self._extract_foreign_languages(text),
            "extraction_mode": mode
        }

    def _find_full_name(self, text: str) -> str:
        """Nhận diện họ và tên ứng viên từ các dòng đầu CV."""
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        for line in lines[:15]:
            # Bỏ qua dòng có email hoặc số điện thoại
            if '@' in line or re.search(r'\d', line):
                continue

            words = line.split()
            if not (2 <= len(words) <= 5):
                continue

            # Bỏ qua nếu dòng chứa các từ khóa tiêu đề chung
            line_lower = line.lower()
            if any(ignore in line_lower for ignore in NAME_IGNORE_KEYWORDS):
                continue

            # Kiểm tra xem có viết hoa chữ cái đầu hoặc viết hoa toàn bộ không
            if line.isupper() or line.istitle() or any(w[0].isupper() for w in words if w):
                # Làm sạch ký tự lạ
                clean_name = re.sub(r'[^a-zA-ZÀ-ỹ\s]', '', line).strip()
                if clean_name:
                    return clean_name.title()

        return "Unknown Candidate"

    def _extract_emails(self, text: str) -> List[str]:
        if not text:
            return []
        pattern = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', re.IGNORECASE)
        return list(set(pattern.findall(text)))

    def _find_phone_number(self, text: str) -> Optional[str]:
        patterns = [
            r'\b((\+84|0)[3|5|7|8|9]([\s\.-]?\d{2,4}){2,3})\b',
            r'\b(\(?[1-9]\d{2}\)?[\s\.-]?\d{3}[\s\.-]?\d{4})\b',
            r'\b([1-9]\d{2}[\s\.-]\d{4})\b'
        ]
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                phone_str = match.group(0)
                if phone_str.startswith('+'):
                    return '+' + re.sub(r'\D', '', phone_str[1:])
                return re.sub(r'\D', '', phone_str)
        return None

    def _find_years_experience(self, text: str) -> float:
        """Ước tính số năm kinh nghiệm từ từ khóa tường minh và các khoảng thời gian đi làm."""
        keyword_years = []

        # 1. Cụm từ rõ ràng: "3 years experience", "2 năm kinh nghiệm", "6 tháng".
        #    Negative lookahead loại tuổi ("22 years old", "22 tuổi") khỏi bộ đếm.
        pattern = re.compile(
            r'(\d{1,2})(?:\.\d)?\+?\s*(year|năm|nam|yr|tháng|month)s?\b(?!\s*(?:old|tuổi|age))',
            re.IGNORECASE,
        )
        for match in pattern.finditer(text):
            num = float(match.group(1))
            unit = match.group(2).lower()
            if 'tháng' in unit or 'month' in unit:
                keyword_years.append(round(num / 12.0, 1))
            elif num <= 30:
                keyword_years.append(num)

        # 2. Cộng dồn khoảng làm việc sau khi gộp đoạn chồng lấn, bỏ khối học vấn.
        intervals = []
        current_year = datetime.now().year
        for line in text.splitlines():
            if not line.strip():
                continue
            line_lower = line.lower()
            if any(kw in line_lower for kw in EDUCATION_LINE_KEYWORDS):
                continue
            for m in YEAR_RANGE_PATTERN.finditer(line):
                start = int(m.group(1))
                end_token = m.group(2).lower()
                if re.match(r'^(present|nay|now|this\s*time|hiện\s*tại)', end_token):
                    end = current_year
                else:
                    end = int(end_token)
                if 0 < end - start <= 15:
                    intervals.append((start, end))

        merged_total = 0.0
        if intervals:
            intervals.sort()
            cur_start, cur_end = intervals[0]
            for start, end in intervals[1:]:
                if start <= cur_end:
                    cur_end = max(cur_end, end)
                else:
                    merged_total += cur_end - cur_start
                    cur_start, cur_end = start, end
            merged_total += cur_end - cur_start

        candidates = keyword_years + ([float(merged_total)] if merged_total > 0 else [])
        return max(candidates) if candidates else 0.0

    def _extract_education_level(self, text: str) -> str:
        """Nhận diện bằng cấp, trả về bậc cao nhất xuất hiện trong CV."""
        text_lower = text.lower()

        # Hỗ trợ viết tắt có/không dấu chấm (Ph.D., B.Sc., M.S.) và không dùng chức danh
        # "engineer" làm bằng cấp để tránh nhầm "Software Engineer" thành Bachelor.
        edu_patterns = {
            'Phd': [
                r'\bph\.?\s?d\.?\b', r'\bdoctorate\b',
                r'\bdoctor\s+of\s+philosophy\b', r'\btiến\s*sĩ\b'
            ],
            'Master': [
                r'\bmasters?(?:\'s)?\b', r'\bthạc\s*sĩ\b', r'\bmba\b',
                r'\bm\.?\s?sc\.?\b', r'\bm\.?\s?tech\.?\b', r'\bm\.\s?s\.'
            ],
            'Bachelor': [
                r'\bbachelors?(?:\'s)?\b', r'\bcử\s*nhân\b', r'\bkỹ\s*sư\b',
                r'\bb\.?\s?sc\.?\b', r'\bb\.?\s?tech\.?\b', r'\bb\.\s?s\.',
                r'\bđại\s*học\b', r'\buniversity\b'
            ],
            'Associate': [
                r'\bassociate\b', r'\bcao\s*đẳng\b', r'\bvocational\b', r'\bcollege\b'
            ]
        }

        # Duyệt từ bậc cao nhất xuống thấp nhất
        for level, patterns in edu_patterns.items():
            for pattern in patterns:
                if re.search(pattern, text_lower):
                    return level

        # Fallback dùng Embedding nếu văn bản phức tạp
        try:
            lines = [s.strip() for s in text.split('\n') if 4 < len(s.strip()) < 80]
            if lines and self._embedder is not None:
                from sentence_transformers import util
                emb_edu = self.embedder.encode(EDUCATION_LEVELS, convert_to_tensor=True)
                emb_lines = self.embedder.encode(lines[:20], convert_to_tensor=True)
                cos_scores = util.cos_sim(emb_lines, emb_edu)
                max_score = cos_scores.max()
                if max_score > 0.65:
                    idx = (cos_scores == max_score).nonzero(as_tuple=True)[1][0]
                    return EDUCATION_LEVELS[idx.item()].title()
        except Exception as e:
            logger.warning(f"Embedding education fallback failed: {e}")

        return "None"

    def _extract_foreign_languages(self, text: str) -> List[str]:
        """Phát hiện ngoại ngữ kèm chứng chỉ (IELTS, TOEIC, JLPT, HSK...) hoặc tên ngôn ngữ đứng riêng."""
        if not text:
            return []

        results: Dict[str, str] = {}
        has_lang_context = bool(re.search(
            r'language|ngôn ngữ|ngoại ngữ|communication|certificate|chứng chỉ',
            text, re.IGNORECASE
        ))

        for lang, patterns in LANGUAGE_SYNONYMS.items():
            match = None
            for pattern in patterns:
                match = re.search(pattern, text, re.IGNORECASE)
                if match:
                    break
            if not match:
                continue

            # 1) Chứng chỉ / mức độ nằm ngay sau tên ngôn ngữ
            tail = text[match.end(): match.end() + 45]
            cert = re.search(r'\b' + LANGUAGE_LEVEL_PATTERN + r'\b', tail, re.IGNORECASE)
            if cert:
                results[lang] = f"{lang} ({self._format_language_level(cert.group(1))})"
                continue

            # 2) Tên ngôn ngữ đứng riêng trong dòng ngắn (mục kỹ năng) hoặc có ngữ cảnh ngoại ngữ
            line_start = text.rfind('\n', 0, match.start()) + 1
            line_end = text.find('\n', match.end())
            if line_end == -1:
                line_end = len(text)
            line = text[line_start:line_end].strip()

            if has_lang_context or (len(line) <= 60 and len(line.split()) <= 8):
                results[lang] = lang

        # 3) Chứng chỉ đứng một mình vẫn suy ra được ngôn ngữ tương ứng
        cert_to_lang = [
            (r'IELTS\s*[\d.]{1,4}|TOEIC\s*\d{3,4}|TOEFL(?:\s*iBT)?\s*[\d.]{1,3}', 'English'),
            (r'JLPT\s*N?[1-5]', 'Japanese'),
            (r'HSK\s*[1-6]', 'Chinese'),
            (r'TOPIK\s*[1-6]', 'Korean'),
            (r'DELF\s*[A-C][12]', 'French'),
            (r'Goethe[-\s]?[A-C][12]', 'German'),
        ]
        for pattern, lang in cert_to_lang:
            if lang in results:
                continue
            match = re.search(r'\b(?:' + pattern + r')\b', text, re.IGNORECASE)
            if match:
                results[lang] = f"{lang} ({self._format_language_level(match.group(0))})"

        return sorted(results.values(), key=lambda value: value.lower())

    @staticmethod
    def _format_language_level(raw: str) -> str:
        """Giữ nguyên viết hoa của tên chứng chỉ (IELTS, TOEIC...) thay vì .title() làm sai."""
        text = raw.strip()
        acronym = re.match(r'(ielts|toeic|toefl|jlpt|hsk|topik|delf|goethe)', text, re.IGNORECASE)
        if acronym:
            token = acronym.group(0)
            return token.upper() + text[len(token):]
        return text.title()

    def _extract_all_skills_passive(self, text: str) -> Dict[str, List[str]]:
        """Quét toàn bộ kỹ năng kỹ thuật có trong CV kèm tra cứu từ đồng nghĩa (Aliases)."""
        found_skills = {}
        text_lower = text.lower()

        def scan_list(category: str, items: List[str]):
            found = set()
            for item in items:
                item_esc = re.escape(item.lower())
                pattern = r'(?:^|\s|[,\.\-\/\(\)])' + item_esc + r'(?:$|\s|[,\.\-\/\(\)])'
                if re.search(pattern, text_lower):
                    # Chuẩn hóa nếu là alias (vd: k8s -> kubernetes)
                    normalized_name = SKILL_ALIASES.get(item.lower(), item).title()
                    found.add(normalized_name)
            if found:
                found_skills[category] = sorted(list(found))

        scan_list('Languages', LANGUAGES)
        scan_list('Frameworks', FRAMEWORKS)
        scan_list('Databases', DATABASES)
        scan_list('AI & Data', AI_ML_SKILLS)
        scan_list('DevOps & Cloud', DEVOPS_TOOLS)

        return found_skills
